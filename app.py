import os
import base64
from io import BytesIO
from pathlib import Path

import streamlit as st
import fitz  # PyMuPDF
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook
from PIL import Image
from groq import Groq

APP_NAME = "StudyMate AI"
TEXT_MODEL_DEFAULT = "llama-3.3-70b-versatile"
VISION_MODEL_DEFAULT = "meta-llama/llama-4-scout-17b-16e-instruct"

st.set_page_config(page_title=APP_NAME, page_icon="📚", layout="wide")

st.title("📚 StudyMate AI")
st.caption("Upload your study materials and ask questions about them.")

st.info(
    "Supported: PDF, DOCX, PPTX, XLSX, TXT, PNG, JPG and JPEG. "
    "Text is extracted from documents; image files are analyzed with a Groq vision model."
)

def get_secret(name: str, default: str = "") -> str:
    """Read a Streamlit secret first, then an environment variable."""
    try:
        value = st.secrets.get(name, "")
        if value:
            return str(value)
    except Exception:
        pass
    return os.getenv(name, default)

def extract_pdf(data: bytes) -> str:
    text_parts = []
    with fitz.open(stream=data, filetype="pdf") as pdf:
        for page_num, page in enumerate(pdf, start=1):
            page_text = page.get_text("text").strip()
            if page_text:
                text_parts.append(f"[Page {page_num}]\n{page_text}")
    return "\n\n".join(text_parts)

def extract_docx(data: bytes) -> str:
    doc = Document(BytesIO(data))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table_num, table in enumerate(doc.tables, start=1):
        parts.append(f"[Table {table_num}]")
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)

def extract_pptx(data: bytes) -> str:
    presentation = Presentation(BytesIO(data))
    parts = []
    for slide_num, slide in enumerate(presentation.slides, start=1):
        slide_parts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_parts.append(shape.text.strip())
            if shape.has_table:
                for row in shape.table.rows:
                    slide_parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        if slide_parts:
            parts.append(f"[Slide {slide_num}]\n" + "\n".join(slide_parts))
    return "\n\n".join(parts)

def extract_xlsx(data: bytes) -> str:
    workbook = load_workbook(BytesIO(data), data_only=True, read_only=True)
    parts = []
    for sheet in workbook.worksheets:
        parts.append(f"[Sheet: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if any(value.strip() for value in values):
                parts.append(" | ".join(values))
    workbook.close()
    return "\n".join(parts)

def extract_txt(data: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")

def extract_document(uploaded_file):
    data = uploaded_file.getvalue()
    suffix = Path(uploaded_file.name).suffix.lower()
    if suffix == ".pdf":
        return {"name": uploaded_file.name, "kind": "text", "text": extract_pdf(data)}
    if suffix == ".docx":
        return {"name": uploaded_file.name, "kind": "text", "text": extract_docx(data)}
    if suffix == ".pptx":
        return {"name": uploaded_file.name, "kind": "text", "text": extract_pptx(data)}
    if suffix == ".xlsx":
        return {"name": uploaded_file.name, "kind": "text", "text": extract_xlsx(data)}
    if suffix == ".txt":
        return {"name": uploaded_file.name, "kind": "text", "text": extract_txt(data)}
    if suffix in (".png", ".jpg", ".jpeg"):
        # Verify that the uploaded file is a readable image before sending it to the vision model.
        image = Image.open(BytesIO(data))
        image.verify()
        return {"name": uploaded_file.name, "kind": "image", "bytes": data, "mime": "image/png" if suffix == ".png" else "image/jpeg"}
    raise ValueError(f"Unsupported file type: {suffix}")

def split_into_chunks(text: str, chunk_size: int = 9000):
    text = text.strip()
    if not text:
        return []
    return [text[i:i + chunk_size] for i in range(0, len(text), chunk_size)]

def relevant_chunks(documents, question: str, max_chunks: int = 6):
    """Simple keyword-based retrieval suitable for a beginner MVP."""
    words = {word.lower().strip(".,?!:;()[]{}\"'") for word in question.split()}
    words = {word for word in words if len(word) > 2}
    candidates = []
    for doc in documents:
        if doc["kind"] != "text":
            continue
        for index, chunk in enumerate(split_into_chunks(doc.get("text", ""))):
            lower_chunk = chunk.lower()
            score = sum(1 for word in words if word in lower_chunk)
            candidates.append((score, doc["name"], index + 1, chunk))
    candidates.sort(key=lambda item: item[0], reverse=True)
    chosen = [item for item in candidates[:max_chunks] if item[0] > 0]
    if not chosen:
        chosen = candidates[:min(3, len(candidates))]
    return [
        f"[Source: {name}, text section {chunk_num}]\n{chunk}"
        for _, name, chunk_num, chunk in chosen
    ]

def get_client():
    api_key = get_secret("GROQ_API_KEY")
    if not api_key:
        return None
    return Groq(api_key=api_key)

def ask_text_model(client, question: str, context: str, model: str):
    system_prompt = (
        "You are StudyMate AI, a helpful study assistant. Answer in clear, student-friendly language. "
        "Use the provided document excerpts as the primary source. If the document does not define a word, "
        "you may explain its general meaning, but clearly say when the definition is general knowledge "
        "rather than something explicitly stated in the document. Never pretend that missing information "
        "appears in the uploaded material. Mention the source filename when useful. If the user asks for "
        "a summary, summarize the available excerpts. If the context is insufficient, say so and ask for "
        "a more specific question or a relevant page/section."
    )
    response = client.chat.completions.create(
        model=model,
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": f"DOCUMENT EXCERPTS:\n{context or '[No readable text found in the uploaded documents.]'}\n\nQUESTION:\n{question}"},
        ],
        temperature=0.2,
        max_tokens=1800,
    )
    return response.choices[0].message.content or "I couldn't generate an answer. Please try again."

def ask_vision_model(client, question: str, image_docs, model: str):
    content = [{
        "type": "text",
        "text": (
            "Answer the student's question using these uploaded image(s). Read visible text when possible, "
            "describe relevant diagrams or tables, and be honest if something is unclear. "
            "If the question asks about a word's meaning, explain it even if the image does not define it, "
            "and distinguish general knowledge from text visible in the image.\n\n"
            f"Student question: {question}"
        )
    }]
    for doc in image_docs[:4]:
        encoded = base64.b64encode(doc["bytes"]).decode("utf-8")
        content.append({"type": "text", "text": f"Uploaded image: {doc['name']}"})
        content.append({
            "type": "image_url",
            "image_url": {"url": f"data:{doc['mime']};base64,{encoded}"}
        })
    response = client.chat.completions.create(
        model=model,
        messages=[
            {"role": "system", "content": "You are a careful, student-friendly study assistant. Do not invent unreadable text."},
            {"role": "user", "content": content},
        ],
        temperature=0.2,
        max_tokens=1800,
    )
    return response.choices[0].message.content or "I couldn't analyze the image. Please try again."

with st.sidebar:
    st.header("⚙️ Settings")
    text_model = st.text_input("Groq text model", value=get_secret("GROQ_TEXT_MODEL", TEXT_MODEL_DEFAULT))
    vision_model = st.text_input("Groq vision model", value=get_secret("GROQ_VISION_MODEL", VISION_MODEL_DEFAULT))
    st.caption("If Groq retires a model, replace its name with a currently available model in your Groq console.")
    st.divider()
    st.markdown("**Privacy tip:** Upload only documents you are allowed to share with an online AI service.")

uploaded_files = st.file_uploader(
    "Upload your study materials",
    type=["pdf", "docx", "pptx", "xlsx", "txt", "png", "jpg", "jpeg"],
    accept_multiple_files=True,
    help="You can select more than one file at a time."
)

if "documents" not in st.session_state:
    st.session_state.documents = []
if "processed_names" not in st.session_state:
    st.session_state.processed_names = set()
if "chat" not in st.session_state:
    st.session_state.chat = []

col1, col2 = st.columns([1, 1])
with col1:
    process_clicked = st.button("📥 Process uploaded files", type="primary", use_container_width=True)
with col2:
    clear_clicked = st.button("🗑️ Clear documents and chat", use_container_width=True)

if clear_clicked:
    st.session_state.documents = []
    st.session_state.processed_names = set()
    st.session_state.chat = []
    st.rerun()

if process_clicked:
    if not uploaded_files:
        st.warning("Please upload at least one file first.")
    else:
        new_docs = []
        errors = []
        progress = st.progress(0)
        for i, uploaded in enumerate(uploaded_files):
            try:
                doc = extract_document(uploaded)
                if doc["kind"] == "text" and not doc["text"].strip():
                    errors.append(f"{uploaded.name}: no selectable text found. If this is a scanned PDF, try uploading clear page images instead.")
                else:
                    new_docs.append(doc)
            except Exception as exc:
                errors.append(f"{uploaded.name}: {exc}")
            progress.progress((i + 1) / len(uploaded_files))
        if new_docs:
            # Replace prior documents with the currently selected processed set.
            st.session_state.documents = new_docs
            st.session_state.processed_names = {doc["name"] for doc in new_docs}
            st.success(f"Processed {len(new_docs)} file(s). You can now ask questions below.")
        for error in errors:
            st.warning(error)

if st.session_state.documents:
    st.subheader("📄 Loaded materials")
    for doc in st.session_state.documents:
        if doc["kind"] == "text":
            length = len(doc.get("text", ""))
            st.write(f"**{doc['name']}** — text document, {length:,} characters")
        else:
            st.write(f"**{doc['name']}** — image")
    st.caption("The current session keeps your uploaded content in memory; clearing the session or refreshing may remove it.")

    st.subheader("💬 Ask your question")
    question = st.text_area(
        "What would you like to know?",
        placeholder="Example: Summarize Chapter 2, explain 'indigenous', or make 10 MCQs from these notes.",
        height=100,
    )
    ask_clicked = st.button("✨ Get answer", type="primary")

    if ask_clicked:
        if not question.strip():
            st.warning("Type a question first.")
        else:
            client = get_client()
            if client is None:
                st.error("Groq API key not found. Add GROQ_API_KEY to Streamlit Secrets or your local environment.")
            else:
                with st.spinner("Reading your materials and preparing an answer..."):
                    try:
                        text_docs = [doc for doc in st.session_state.documents if doc["kind"] == "text"]
                        image_docs = [doc for doc in st.session_state.documents if doc["kind"] == "image"]
                        chunks = relevant_chunks(text_docs, question)
                        context = "\n\n---\n\n".join(chunks)
                        if image_docs and (not context or any(term in question.lower() for term in ["image", "picture", "photo", "diagram", "chart", "shown", "this"])):
                            answer = ask_vision_model(client, question, image_docs, vision_model)
                        elif context:
                            answer = ask_text_model(client, question, context, text_model)
                        elif image_docs:
                            answer = ask_vision_model(client, question, image_docs, vision_model)
                        else:
                            st.error("No readable document content was found. Try another file or a more text-based PDF.")
                            answer = None
                        if answer:
                            st.session_state.chat.append({"question": question.strip(), "answer": answer})
                    except Exception as exc:
                        st.error(
                            "The request failed. Check your API key, model names, internet connection, "
                            f"and Groq usage limits. Technical detail: {exc}"
                        )

    if st.session_state.chat:
        st.divider()
        st.subheader("📝 Conversation")
        for item in reversed(st.session_state.chat):
            with st.chat_message("user"):
                st.markdown(item["question"])
            with st.chat_message("assistant"):
                st.markdown(item["answer"])
else:
    st.write("👆 Upload one or more files and click **Process uploaded files** to begin.")

st.divider()
st.caption("StudyMate AI is a learning aid. Always verify important facts against your original syllabus or notes.")
